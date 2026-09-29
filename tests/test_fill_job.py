"""A whole filling job, end to end: the real agent loop and tools, a mocked Claude."""

import openpyxl
import pytest

from coding_agent import config, memory, session
from excel_filler import agent

from .helpers import FakeClaude, ScriptedUI, make_pdf


@pytest.fixture
def folder(tmp_path, monkeypatch):
    monkeypatch.setattr(session, "MEMORY_HOME", tmp_path / "memory-home")
    monkeypatch.setattr(config, "MEMORY_UPDATES", False)
    monkeypatch.setattr(memory, "MEMORY_UPDATES", False)
    import coding_agent.tools.documents as documents
    monkeypatch.setattr(documents, "BACKUP_HOME", tmp_path / "backups")
    f = tmp_path / "invoices"
    f.mkdir()
    make_pdf(f / "invoice.pdf", ["Invoice INV-31: Sensors 12 x 45.50 EUR", "Total 546.00 EUR"])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Costs"
    ws.append(["Item", "Qty", "Unit price", "Total"])
    ws["D2"] = "=B2*C2"
    wb.save(f / "costs.xlsx")
    return f


def test_fill_job_end_to_end(folder, monkeypatch):
    fake = FakeClaude([
        ([("read_pdf", {"path": "invoice.pdf", "mode": "text"})], "tool_use"),        # tries the PDF first
        ([("read_excel", {"path": "costs.xlsx"})], "tool_use"),
        ([("read_pdf", {"path": "invoice.pdf", "pages": "1"})], "tool_use"),
        ([("edit_excel", {"path": "costs.xlsx", "changes": [
            {"sheet": "Costs", "cell": "A2", "value": "Sensors"},
            {"sheet": "Costs", "cell": "B2", "value": 12},
            {"sheet": "Costs", "cell": "C2", "value": 45.5}]})], "tool_use"),
        ([("text", "Filled row 2 from invoice.pdf page 1.")], "end_turn"),
    ])
    monkeypatch.setattr(session, "_get_client", fake.client)
    ui = ScriptedUI(answers=["yes"])
    agent.open_folder(folder, ui)
    assert agent.fill("costs.xlsx", ["invoice.pdf"], notes="amounts in EUR")

    first = fake.requests[0]
    assert sorted(t["name"] for t in first["tools"]) == sorted(agent.TOOLS)
    assert first["system"].startswith("You fill Excel workbooks") and str(folder.resolve()) in first["system"]
    instruction = first["messages"][0]["content"][-1]["text"]
    assert "Fill the Excel workbook costs.xlsx" in instruction and "- invoice.pdf" in instruction and "amounts in EUR" in instruction
    assert "read_excel first" in fake.tool_results(1)["read"]["content"]  # the spreadsheet-first rule
    doc = fake.tool_results(3)["read"]["content"]
    assert [b["type"] for b in doc] == ["text", "document"]

    ws = openpyxl.load_workbook(folder / "costs.xlsx")["Costs"]
    assert (ws["A2"].value, ws["B2"].value, ws["C2"].value, ws["D2"].value) == ("Sensors", 12, 45.5, "=B2*C2")
    assert ("confirm", "Apply these cell changes to costs.xlsx?") in ui.events
    assert ("text", "Filled row 2 from invoice.pdf page 1.") in ui.events


def test_documents_in_another_folder_are_read_not_written(folder, tmp_path, monkeypatch):
    other = tmp_path / "scans"
    other.mkdir()
    make_pdf(other / "inv-9.pdf", ["Brackets 8 x 12.00 EUR"])
    pdf = (other / "inv-9.pdf").resolve().as_posix()
    fake = FakeClaude([
        ([("read_excel", {"path": "costs.xlsx"})], "tool_use"),
        ([("read_pdf", {"path": pdf, "mode": "text"}), ("edit_excel", {"path": str(other / "copy.xlsx"), "changes": [{"cell": "A1", "value": 1}]})], "tool_use"),
        ([("text", "done")], "end_turn"),
    ])
    monkeypatch.setattr(session, "_get_client", fake.client)
    agent.open_folder(folder, ScriptedUI(answers=["yes"]))
    session.add_read_folder(other)
    agent.fill("costs.xlsx", [pdf])
    first = fake.requests[0]["messages"][0]["content"]
    assert any(b["text"].startswith("<read_only_folders>") for b in first)
    results = fake.tool_results(2)
    assert "Brackets 8 x 12.00 EUR" in results["read"]["content"]
    assert results["edit"]["is_error"] and not (other / "copy.xlsx").exists()


def test_tools_outside_the_excel_set_are_not_available(folder, monkeypatch):
    fake = FakeClaude([
        ([("run_python", {"code": "print(1)"}), ("delete_file", {"path": "costs.xlsx"})], "tool_use"),
        ([("text", "ok")], "end_turn"),
    ])
    monkeypatch.setattr(session, "_get_client", fake.client)
    agent.open_folder(folder, ScriptedUI())
    agent.fill("costs.xlsx", [])
    results = fake.tool_results(1)
    assert results["run"]["is_error"] and results["delete"]["is_error"]
    assert (folder / "costs.xlsx").exists()
