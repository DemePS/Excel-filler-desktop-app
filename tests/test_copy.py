"""Working on a copy: the original workbook is never changed."""

import openpyxl
import pytest
from fastapi.testclient import TestClient

from coding_agent import backups, config, memory, session
from excel_filler.desktop import copies
from excel_filler.desktop.server import create_app

from .helpers import FakeClaude, make_pdf

TOKEN = "t0ken"


def test_the_copy_is_named_next_to_the_original_and_never_overwrites(tmp_path):
    (tmp_path / "costs.xlsx").write_bytes(b"original")
    assert copies.make_copy(tmp_path, "costs.xlsx") == "costs (copy).xlsx"
    assert (tmp_path / "costs (copy).xlsx").read_bytes() == b"original"
    (tmp_path / "costs (copy).xlsx").write_bytes(b"edited copy")
    assert copies.make_copy(tmp_path, "costs.xlsx") == "costs (copy 2).xlsx"          # the first copy is kept
    assert (tmp_path / "costs (copy).xlsx").read_bytes() == b"edited copy"
    assert copies.make_copy(tmp_path, "costs (copy).xlsx") == "costs (copy 3).xlsx"   # a copy of a copy: no "(copy) (copy)"
    (tmp_path / "macros.xlsm").write_bytes(b"m")
    assert copies.make_copy(tmp_path, "macros.xlsm") == "macros (copy).xlsm"           # the extension is kept


def test_only_a_workbook_of_the_folder_can_be_copied(tmp_path):
    (tmp_path / "notes.txt").write_text("x")
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub" / "a.xlsx").write_bytes(b"a")
    for bad in ("notes.txt", "missing.xlsx", "sub/a.xlsx", "../a.xlsx", "/etc/passwd"):
        with pytest.raises(ValueError):
            copies.make_copy(tmp_path, bad)
    assert sorted(p.name for p in tmp_path.iterdir()) == ["notes.txt", "sub"]          # nothing was created


@pytest.fixture
def job(tmp_path, monkeypatch):
    monkeypatch.setattr(session, "MEMORY_HOME", tmp_path / "memory-home")
    monkeypatch.setattr(config, "MEMORY_UPDATES", False)
    monkeypatch.setattr(memory, "MEMORY_UPDATES", False)
    monkeypatch.setenv("ANTHROPIC_FOUNDRY_ENDPOINT", "https://x.services.ai.azure.com/anthropic")
    monkeypatch.setattr(backups, "BACKUP_HOME", tmp_path / "backups")
    folder = tmp_path / "job"
    folder.mkdir()
    make_pdf(folder / "invoice.pdf", ["Sensors 12 x 45.50 EUR"])
    wb = openpyxl.Workbook()
    wb.active.title = "Costs"
    wb.active.append(["Item", "Qty"])
    wb.save(folder / "costs.xlsx")
    return TestClient(create_app(TOKEN), headers={"x-token": TOKEN}), folder


def run_job(client, folder, monkeypatch, copy):
    fake = FakeClaude([
        ([("read_excel", {"path": "costs (copy).xlsx" if copy else "costs.xlsx"})], "tool_use"),
        ([("edit_excel", {"path": "costs (copy).xlsx" if copy else "costs.xlsx", "changes": [
            {"sheet": "Costs", "cell": "A2", "value": "Sensors"}, {"sheet": "Costs", "cell": "B2", "value": 12}]})], "tool_use"),
        ([("text", "Filled row 2.")], "end_turn"),
    ])
    monkeypatch.setattr(session, "_get_client", fake.client)
    client.post("/api/folder", json={"path": str(folder)})
    with client.websocket_connect(f"/ws?token={TOKEN}") as ws:
        ws.receive_json()
        started = client.post("/api/job", json={"workbook": "costs.xlsx", "documents": ["invoice.pdf"], "on_copy": copy}).json()
        events = []
        while True:
            event = ws.receive_json()
            events.append(event)
            if event["type"] == "confirm":
                client.post("/api/answer", json={"id": event["id"], "value": "yes"})
            if event["type"] == "busy" and event["busy"] is False:
                break
    return started, events


def test_a_job_on_a_copy_fills_the_copy_and_leaves_the_original_alone(job, monkeypatch):
    client, folder = job
    before = (folder / "costs.xlsx").read_bytes()
    started, events = run_job(client, folder, monkeypatch, copy=True)
    assert started == {"started": True, "workbook": "costs (copy).xlsx"}
    assert (folder / "costs.xlsx").read_bytes() == before                              # the original: not one byte changed
    ws = openpyxl.load_workbook(folder / "costs (copy).xlsx")["Costs"]
    assert (ws["A2"].value, ws["B2"].value) == ("Sensors", 12)
    request = next(e for e in events if e["type"] == "request")
    assert request["text"] == "Fill costs (copy).xlsx from invoice.pdf"                 # the feed says which file
    listings = [e for e in events if e["type"] == "listing"]
    assert listings and "costs (copy).xlsx" in listings[0]["workbooks"] and listings[0]["workbook"] == "costs (copy).xlsx"


def test_without_the_copy_the_workbook_itself_is_filled(job, monkeypatch):
    client, folder = job
    started, _ = run_job(client, folder, monkeypatch, copy=False)
    assert started == {"started": True, "workbook": "costs.xlsx"}
    assert not (folder / "costs (copy).xlsx").exists()
    assert openpyxl.load_workbook(folder / "costs.xlsx")["Costs"]["B2"].value == 12


def test_a_missing_workbook_is_refused_and_no_copy_is_made(job):
    client, folder = job
    client.post("/api/folder", json={"path": str(folder)})
    r = client.post("/api/job", json={"workbook": "absent.xlsx", "documents": [], "on_copy": True})
    assert r.status_code == 400
    assert sorted(p.name for p in folder.iterdir()) == ["costs.xlsx", "invoice.pdf"]
