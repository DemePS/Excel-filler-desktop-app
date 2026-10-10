"""Guardrails that hold whatever the model is told: formulas, and the one workbook of the job."""

import openpyxl
import pytest
from fastapi.testclient import TestClient

from coding_agent import backups, config, memory, session, state
from excel_filler.desktop.server import create_app

from .helpers import FakeClaude

TOKEN = "t0ken"


@pytest.fixture
def world(tmp_path, monkeypatch):
    monkeypatch.setattr(session, "MEMORY_HOME", tmp_path / "memory-home")
    monkeypatch.setattr(config, "MEMORY_UPDATES", False)
    monkeypatch.setattr(memory, "MEMORY_UPDATES", False)
    monkeypatch.setattr(backups, "BACKUP_HOME", tmp_path / "backups")
    monkeypatch.setenv("ANTHROPIC_FOUNDRY_ENDPOINT", "https://x.services.ai.azure.com/anthropic")
    folder = tmp_path / "job"
    folder.mkdir()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Journal"
    ws.append(["HT", "TVA", "Contrôle"])
    ws["C2"] = '=IF(A2="","",IF(ROUND(A2*0.2,2)=B2,"OK","ECART"))'
    wb.save(folder / "journal.xlsx")
    return TestClient(create_app(TOKEN), headers={"x-token": TOKEN}), folder


def run(client, folder, monkeypatch, script):
    fake = FakeClaude(script)
    monkeypatch.setattr(session, "_get_client", fake.client)
    client.post("/api/folder", json={"path": str(folder)})
    with client.websocket_connect(f"/ws?token={TOKEN}") as ws:
        ws.receive_json()
        assert client.post("/api/job", json={"workbook": "journal.xlsx", "documents": []}).json()["started"] is True
        while True:
            event = ws.receive_json()
            if event["type"] == "confirm":
                client.post("/api/answer", json={"id": event["id"], "value": "yes"})
            if event["type"] == "busy" and event["busy"] is False:
                break
    return fake


def test_a_formula_can_be_neither_written_nor_changed_nor_cleared(world, monkeypatch):
    client, folder = world
    fake = run(client, folder, monkeypatch, [
        ([("read_excel", {"path": "journal.xlsx"})], "tool_use"),
        ([("edit_excel", {"path": "journal.xlsx", "changes": [{"sheet": "Journal", "cell": "C2", "value": "OK"}]})], "tool_use"),       # replace
        ([("edit_excel", {"path": "journal.xlsx", "changes": [{"sheet": "Journal", "cell": "C2", "value": None}]})], "tool_use"),       # clear
        ([("edit_excel", {"path": "journal.xlsx", "changes": [{"sheet": "Journal", "cell": "A2", "value": "=1+1"}]})], "tool_use"),     # write one
        ([("text", "Done.")], "end_turn"),
    ])
    results = [fake.tool_results(i)["edit"]["content"] for i in (2, 3, 4)]
    assert all("formula" in str(r).lower() for r in results), results
    ws = openpyxl.load_workbook(folder / "journal.xlsx")["Journal"]
    assert str(ws["C2"].value).startswith("=IF(") and ws["A2"].value is None                                                           # untouched


def test_the_job_names_the_one_workbook_it_may_change(world, monkeypatch):
    client, folder = world
    run(client, folder, monkeypatch, [([("text", "Nothing to do.")], "end_turn")])
    assert state.excel_writable == {(folder / "journal.xlsx").resolve()}
    assert state.excel_protect_formulas is True
