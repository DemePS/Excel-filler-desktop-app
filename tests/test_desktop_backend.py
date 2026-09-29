"""The desktop backend: token protection, and a whole job driven like the window drives it."""

import openpyxl
import pytest
from fastapi.testclient import TestClient

from coding_agent import config, memory, session
from excel_filler.desktop.server import create_app

from .helpers import FakeClaude, make_pdf

TOKEN = "t0ken"


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(session, "MEMORY_HOME", tmp_path / "memory-home")
    monkeypatch.setattr(config, "MEMORY_UPDATES", False)
    monkeypatch.setattr(memory, "MEMORY_UPDATES", False)
    monkeypatch.setenv("ANTHROPIC_FOUNDRY_ENDPOINT", "https://x.services.ai.azure.com/anthropic")
    import coding_agent.tools.documents as documents
    monkeypatch.setattr(documents, "BACKUP_HOME", tmp_path / "backups")
    return TestClient(create_app(TOKEN), headers={"x-token": TOKEN})


@pytest.fixture
def folder(tmp_path):
    f = tmp_path / "job"
    f.mkdir()
    make_pdf(f / "invoice.pdf", ["Sensors 12 x 45.50 EUR"])
    wb = openpyxl.Workbook()
    wb.active.title = "Costs"
    wb.active.append(["Item", "Qty"])
    wb.save(f / "costs.xlsx")
    return f


def test_requests_without_the_token_are_refused(client):
    assert client.get("/api/state", headers={"x-token": "wrong"}).status_code == 403
    assert client.get("/api/state", headers={"x-token": TOKEN, "host": "evil.example"}).status_code == 403
    assert client.get("/api/state").status_code == 200


def test_open_folder_lists_workbooks_and_documents(client, folder):
    r = client.post("/api/folder", json={"path": str(folder)})
    assert r.status_code == 200
    assert r.json()["workbooks"] == ["costs.xlsx"] and r.json()["documents"] == ["invoice.pdf"]
    assert client.post("/api/folder", json={"path": str(folder / "nope")}).status_code == 400


def test_a_whole_job_through_the_api(client, folder, monkeypatch):
    fake = FakeClaude([
        ([("read_excel", {"path": "costs.xlsx"})], "tool_use"),
        ([("edit_excel", {"path": "costs.xlsx", "changes": [
            {"sheet": "Costs", "cell": "A2", "value": "Sensors"}, {"sheet": "Costs", "cell": "B2", "value": 12}]})], "tool_use"),
        ([("text", "Filled row 2 from invoice.pdf page 1.")], "end_turn"),
    ])
    monkeypatch.setattr(session, "_get_client", fake.client)
    client.post("/api/folder", json={"path": str(folder)})
    with client.websocket_connect(f"/ws?token={TOKEN}") as ws:
        hello = ws.receive_json()
        assert hello["type"] == "hello" and hello["workbooks"] == ["costs.xlsx"]
        assert client.post("/api/job", json={"workbook": "costs.xlsx", "documents": ["invoice.pdf"]}).json() == {"started": True}
        events = []
        while True:
            event = ws.receive_json()
            events.append(event)
            if event["type"] == "confirm":
                assert event["question"] == "Apply these cell changes to costs.xlsx?"
                assert client.post("/api/job", json={"workbook": "costs.xlsx"}).status_code == 409  # one job at a time
                client.post("/api/answer", json={"id": event["id"], "value": "yes"})
            if event["type"] == "busy" and event["busy"] is False:
                break
    kinds = [e["type"] for e in events]
    assert kinds[0] == "busy" and "cells" in kinds and "success" in kinds
    assert events[1] == {"type": "request", "text": "Fill costs.xlsx from invoice.pdf"}  # shown first in the feed
    cells = next(e for e in events if e["type"] == "cells")
    assert {"cell": "Costs!B2", "old": "", "new": "12", "format": None} in cells["rows"]
    text = "".join(e["text"] for e in events if e["type"] == "text")
    assert text == "Filled row 2 from invoice.pdf page 1."
    ws_ = openpyxl.load_workbook(folder / "costs.xlsx")["Costs"]
    assert (ws_["A2"].value, ws_["B2"].value) == ("Sensors", 12)


def test_stop_answers_pending_questions_negatively(client, folder, monkeypatch):
    fake = FakeClaude([
        ([("edit_excel", {"path": "costs.xlsx", "changes": [{"cell": "A2", "value": "x"}]})], "tool_use"),
        ([("text", "never")], "end_turn"),
    ])
    monkeypatch.setattr(session, "_get_client", fake.client)
    client.post("/api/folder", json={"path": str(folder)})
    with client.websocket_connect(f"/ws?token={TOKEN}") as ws:
        ws.receive_json()
        client.post("/api/job", json={"workbook": "costs.xlsx"})
        while (event := ws.receive_json())["type"] != "confirm":
            pass
        client.post("/api/stop")
        while not ((event := ws.receive_json())["type"] == "busy" and event["busy"] is False):
            pass
    assert openpyxl.load_workbook(folder / "costs.xlsx")["Costs"]["A2"].value is None


def test_the_real_server_accepts_the_window_websocket(caplog):
    """Through uvicorn itself (the test client above bypasses it): the WebSocket must work."""
    import asyncio
    import json
    import socket
    import threading
    import time

    import uvicorn
    import websockets

    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        port = s.getsockname()[1]
    server = uvicorn.Server(uvicorn.Config(create_app(TOKEN), host="127.0.0.1", port=port, log_level="warning"))
    threading.Thread(target=server.run, daemon=True).start()
    while not server.started:
        time.sleep(0.02)

    async def hello():
        async with websockets.connect(f"ws://127.0.0.1:{port}/ws?token={TOKEN}") as ws:
            return json.loads(await ws.recv())

    async def drop():  # the window closes without a clean goodbye (code 1005, as seen on Windows)
        ws = await websockets.connect(f"ws://127.0.0.1:{port}/ws?token={TOKEN}")
        await ws.recv()
        ws.transport.close()

    try:
        assert asyncio.run(hello())["type"] == "hello"
        asyncio.run(drop())
        time.sleep(0.5)
        import gc
        gc.collect()  # "Task exception was never retrieved" is logged when the task is collected
        time.sleep(0.2)
    finally:
        server.should_exit = True
    problems = [r.getMessage() for r in caplog.records if r.levelname == "ERROR"]
    assert not problems, problems


def test_the_launcher_imports_without_a_console(tmp_path):
    """Started as a windowed app on Windows, sys.stdout/stderr are None: importing must not crash."""
    import subprocess
    import sys
    marker = tmp_path / "ok"
    code = ("import sys; sys.stdout = sys.stderr = None\n"
            "import excel_filler.desktop.app, excel_filler.desktop.server\n"
            "open(sys.argv[1], 'w').write('ok')")
    subprocess.run([sys.executable, "-c", code, str(marker)], check=True)
    assert marker.read_text() == "ok"


def test_the_window_api_exposes_only_its_methods():
    """pywebview publishes every public attribute to the page; the native window must stay private."""
    from excel_filler.desktop.app import WindowApi
    api = WindowApi()
    api._window = object()
    public = [name for name in vars(api) if not name.startswith("_")]
    assert public == []
    assert callable(api.pick_workbook) and callable(api.pick_documents)


def test_open_workbook_opens_its_folder_with_it_selected(client, folder):
    r = client.post("/api/workbook", json={"path": str(folder / "costs.xlsx")})
    assert r.status_code == 200
    assert r.json()["folder"] == str(folder.resolve()) and r.json()["workbook"] == "costs.xlsx"
    old = folder / "old.xls"
    old.write_bytes(b"x")
    assert "saved as .xlsx" in client.post("/api/workbook", json={"path": str(old)}).json()["detail"]
    assert client.post("/api/workbook", json={"path": str(folder / "missing.xlsx")}).status_code == 400


def test_add_documents_from_subfolders_and_other_folders(client, folder, tmp_path):
    client.post("/api/workbook", json={"path": str(folder / "costs.xlsx")})
    (folder / "march").mkdir()
    make_pdf(folder / "march" / "inv-7.pdf", ["x"])
    r = client.post("/api/documents", json={"paths": [str(folder / "march" / "inv-7.pdf")]})
    assert r.status_code == 200 and r.json()["added"] == ["march/inv-7.pdf"]
    assert "march/inv-7.pdf" in client.get("/api/state").json()["documents"]
    other = tmp_path / "scans"
    other.mkdir()
    make_pdf(other / "elsewhere.pdf", ["x"])
    r = client.post("/api/documents", json={"paths": [str(other / "elsewhere.pdf")]})
    absolute = (other / "elsewhere.pdf").resolve().as_posix()
    assert r.status_code == 200 and r.json()["added"] == [absolute]  # from anywhere, by absolute path
    from coding_agent import state
    assert other.resolve() in state.read_roots  # its folder is readable (not writable) by the agent
    (other / "notes.docx").write_bytes(b"x")
    r = client.post("/api/documents", json={"paths": [str(other / "notes.docx")]})
    assert r.status_code == 400 and "notes.docx" in r.json()["detail"]


def test_the_job_log_says_what_the_agent_does(client, folder, monkeypatch, caplog):
    import logging
    caplog.set_level(logging.INFO, logger="excel-filler.job")
    fake = FakeClaude([
        ([("text", "Reading the workbook first."), ("read_excel", {"path": "costs.xlsx"})], "tool_use"),
        ([("read_pdf", {"path": "invoice.pdf", "pages": "1"})], "tool_use"),
        ([("edit_excel", {"path": "costs.xlsx", "changes": [{"sheet": "Costs", "cell": "A2", "value": "Sensors"}]})], "tool_use"),
        ([("text", "Filled A2 from invoice.pdf page 1.")], "end_turn"),
    ])
    monkeypatch.setattr(session, "_get_client", fake.client)
    client.post("/api/folder", json={"path": str(folder)})
    with client.websocket_connect(f"/ws?token={TOKEN}") as ws:
        ws.receive_json()
        client.post("/api/job", json={"workbook": "costs.xlsx", "documents": ["invoice.pdf"]})
        while True:
            event = ws.receive_json()
            if event["type"] == "confirm":
                client.post("/api/answer", json={"id": event["id"], "value": "yes"})
            if event["type"] == "busy" and event["busy"] is False:
                break
    lines = [r.getMessage() for r in caplog.records if r.name == "excel-filler.job"]
    expected = ["Fill costs.xlsx from 1 document(s)", "Job started", "Claude: Reading the workbook first.",
                "Claude reads the workbook costs.xlsx", "Claude reads invoice.pdf (page 1)",
                "Claude prepares changes to costs.xlsx", "Question: Apply these cell changes to costs.xlsx?",
                "You answered: approved", "Done: Modified costs.xlsx (1 cell(s))", "Claude: Filled A2 from invoice.pdf page 1."]
    positions = []
    for line in expected:
        found = [i for i, l in enumerate(lines) if line in l]
        assert found, (line, lines)
        positions.append(found[0])
    assert positions == sorted(positions), lines  # in the order it happened
    assert any(l.startswith("Job finished in") for l in lines)


def test_the_job_log_shows_each_tool_call_and_its_outcome(client, folder, monkeypatch, caplog):
    import logging
    caplog.set_level(logging.INFO, logger="excel-filler.job")
    fake = FakeClaude([
        ([("read_pdf", {"path": "missing.pdf"})], "tool_use"),
        ([("read_excel", {"path": "costs.xlsx"})], "tool_use"),
        ([("text", "Nothing to fill.")], "end_turn"),
    ])
    monkeypatch.setattr(session, "_get_client", fake.client)
    client.post("/api/folder", json={"path": str(folder)})
    with client.websocket_connect(f"/ws?token={TOKEN}") as ws:
        ws.receive_json()
        client.post("/api/job", json={"workbook": "costs.xlsx", "documents": ["invoice.pdf"]})
        events = []
        while not (events and events[-1]["type"] == "busy" and events[-1]["busy"] is False):
            events.append(ws.receive_json())
    results = [e for e in events if e["type"] == "tool_result"]
    assert [(e["name"], e["ok"]) for e in results] == [("read_pdf", False), ("read_excel", True)]
    lines = [r.getMessage() for r in caplog.records if r.name == "excel-filler.job"]
    failed = next(l for l in lines if l.startswith("Tool read_pdf("))
    assert "path='missing.pdf'" in failed and "-> ERROR:" in failed
    assert any(l.startswith("Tool read_excel(path='costs.xlsx') -> OK (") for l in lines), lines


def test_window_errors_are_logged(client, caplog):
    import logging
    caplog.set_level(logging.ERROR, logger="excel-filler.window")
    assert client.post("/api/client-error", json={"message": "TypeError: x is undefined", "stack": "at Change"}).json() == {"ok": True}
    assert client.post("/api/client-error", json={"message": "x"}, headers={"x-token": "wrong"}).status_code == 403
    assert any("Window error: TypeError: x is undefined" in r.getMessage() and "at Change" in r.getMessage()
               for r in caplog.records)


def test_an_event_json_cannot_encode_does_not_break_the_connection():
    import datetime
    from pathlib import Path
    from excel_filler.desktop.server import Desktop
    desktop = Desktop()
    client = TestClient(create_app(TOKEN, desktop))
    with client.websocket_connect(f"/ws?token={TOKEN}") as ws:
        ws.receive_json()
        desktop.ui.emit({"type": "message", "text": "odd", "when": datetime.date(2026, 9, 29), "where": Path("a")})
        desktop.ui.message("next")
        assert ws.receive_json()["when"] == "2026-09-29"
        assert ws.receive_json()["text"] == "next"


def test_job_request_in_plain_words():
    from excel_filler.desktop.server import job_request
    assert job_request("costs.xlsx", ["a.pdf", "/x/y/b.pdf", "c.png"], " excl. VAT ") == \
        "Fill costs.xlsx from a.pdf, b.pdf and c.png\nexcl. VAT"
    assert job_request("costs.xlsx", ["a.pdf"], "") == "Fill costs.xlsx from a.pdf"
